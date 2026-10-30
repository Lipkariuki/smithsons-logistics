"""Simple facade tests; only disposable local PostgreSQL, never production."""
import io
import unittest
from unittest.mock import patch
from contextlib import nullcontext
from datetime import date
from pathlib import Path
from time import perf_counter
import json
from decimal import Decimal
from openpyxl import load_workbook
from sqlalchemy import text
import test_monthly_expense_api as api
import test_monthly_expense_persistence as fixture
import test_saved_monthly_files as saved
from services import monthly_payouts as payout
from services.monthly_expense_persistence import create_draft, get_import, VersionConflict, ValidationBlocked
from services.truck_confirmations import confirm_truck
from services.monthly_statement_documents import create_document, get_document
from routers.monthly_payouts import create_router


@unittest.skipUnless(fixture.URL, 'Explicit disposable PostgreSQL required')
class SimplePayoutTests(unittest.TestCase):
    setUpClass=classmethod(api.ApiTests.setUpClass.__func__)
    tearDownClass=classmethod(api.ApiTests.tearDownClass.__func__)
    tearDown=api.ApiTests.tearDown
    request=api.ApiTests.request
    legacy=api.ApiTests.legacy
    archive=saved.SavedFileTests.archive

    def setUp(self):
        saved.SavedFileTests.setUp(self)
        def connection():
            with self.engine.begin() as conn: yield conn
        self.app.include_router(create_router(connection,api.auth_functions(self.engine)))

    def sync(self,conn,**options):
        latest=payout.latest(conn,fixture.MONTH);raw=payout.raw_latest(conn,fixture.MONTH)
        return payout.synchronize(conn,fixture.MONTH,7,expected_report_id=latest['id'] if latest else None,
            expected_raw_id=raw['id'] if raw else None,**options)

    def test_two_originals_autosave_calculate_and_refresh_idempotently(self):
        with self.engine.begin() as conn:
            before=self.legacy(conn);self.archive(conn)
            first=self.sync(conn,content=fixture.book(),filename='original-net.xlsx')
            row=first['rows'][0]
            self.assertEqual([row[k] for k in ['gross','operating_expenses','commission','net_payable']],list(map(Decimal,['1000','200','56','744'])))
            self.assertEqual(self.sync(conn)['id'],first['id'])
            self.assertEqual(self.legacy(conn),before)
            self.assertEqual(conn.execute(text('SELECT count(*) FROM dhl_monthly_truck_confirmations')).scalar_one(),0)

    def test_get_old_draft_previews_saved_sources_without_writes(self):
        with self.engine.begin() as conn:
            old=create_draft(conn,fixture.book(),fixture.MONTH,'net.xlsx',7);self.archive(conn)
        status,body,_=self.request('GET','/dhl-reports/monthly-payouts/2026-02-01')
        self.assertEqual(status,200,body);self.assertEqual(body['report']['status'],'preview')
        self.assertEqual(body['report']['rows'][0]['gross'],'1000.00')
        with self.engine.begin() as conn:
            self.assertIsNone(get_import(conn,old['id'])['rows'][0]['gross'])
            self.assertEqual(conn.execute(text('SELECT count(*) FROM dhl_monthly_expense_imports')).scalar_one(),1)

    def test_old_alias_rejection_reopens_readonly_and_explicit_review_creates_new_version(self):
        wb=load_workbook(io.BytesIO(self.raw));wb.active['D3']='Offload';out=io.BytesIO();wb.save(out);wb.close()
        legacy_issue=dict(code='unsupported_source_charge',message='Old strict validator rejected Offload',cell=None,blocking=True)
        with self.engine.begin() as conn:
            before=self.legacy(conn);self.archive(conn,out.getvalue())
            with patch('services.monthly_expense_dates.charge_issues',return_value=[legacy_issue]):
                old=self.sync(conn,content=fixture.book(),filename='net.xlsx')
            self.assertIsNone(old['rows'][0]['gross'])
        status,body,_=self.request('GET','/dhl-reports/monthly-payouts/2026-02-01')
        self.assertEqual(status,200,body);self.assertEqual(body['report']['status'],'preview')
        self.assertEqual(body['report']['rows'][0]['gross'],'1000.00')
        with self.engine.begin() as conn:
            self.assertEqual(get_import(conn,old['id']),old)
            self.assertEqual(conn.execute(text('SELECT count(*) FROM dhl_monthly_expense_imports')).scalar_one(),1)
            self.assertEqual(conn.execute(text('SELECT count(*) FROM dhl_monthly_truck_confirmations')).scalar_one(),0)
            revised=self.sync(conn)
            self.assertEqual(revised['version'],2)
            self.assertIn('Known COST aliases revalidated',revised['decisions']['reason'])
            self.assertEqual(revised['rows'][0]['net_payable'],Decimal('744.00'))
            self.assertEqual(self.sync(conn)['id'],revised['id'])
            confirm_truck(conn,revised['rows'][0]['id'],7,'Synthetic alias review only')
            doc=create_document(conn,revised['rows'][0]['id'],7)
            self.assertEqual(bytes(get_document(conn,doc['id'])['pdf_content']),bytes(doc['pdf_content']))
            self.assertEqual(self.legacy(conn),before)

    def test_unknown_cost_stays_blocked_without_a_spurious_alias_revision(self):
        wb=load_workbook(io.BytesIO(self.raw));wb.active['D3']='Dist / Offload';out=io.BytesIO();wb.save(out);wb.close()
        with self.engine.begin() as conn:
            self.archive(conn,out.getvalue());old=self.sync(conn,content=fixture.book(),filename='net.xlsx')
            _,current=payout.state(conn,fixture.MONTH)
            self.assertEqual(current['id'],old['id'])
            self.assertIsNone(current['rows'][0]['gross'])
            with self.assertRaises(ValidationBlocked):confirm_truck(conn,current['rows'][0]['id'],7,'Cannot guess COST')

    def test_replacement_preserves_pdf_and_never_accumulates_expenses(self):
        with self.engine.begin() as conn:
            self.archive(conn);first=self.sync(conn,content=fixture.book(),filename='net.xlsx')
            approval=confirm_truck(conn,first['rows'][0]['id'],7,'Synthetic independent review')
            old=create_document(conn,first['rows'][0]['id'],7)
            second=self.sync(conn,content=fixture.book(200),filename='corrected.xlsx')
            self.assertEqual(second['version'],2);self.assertEqual(second['rows'][0]['operating_expenses'],Decimal('300.00'))
            self.assertEqual(second['rows'][0]['net_payable'],Decimal('651.00'))
            self.assertEqual(bytes(get_document(conn,old['id'])['pdf_content']),bytes(old['pdf_content']))
            self.assertEqual(conn.execute(text('SELECT status FROM dhl_monthly_truck_confirmations WHERE id=:id'),{'id':approval['id']}).scalar_one(),'current')

    def test_unknown_categories_need_explicit_classification(self):
        with self.engine.begin() as conn:
            self.archive(conn);first=self.sync(conn,content=fixture.book(extra=True),filename='tyres.xlsx')
            self.assertIsNone(first['rows'][0]['net_payable'])
            with self.assertRaises(ValidationBlocked):confirm_truck(conn,first['rows'][0]['id'],7,'Must not guess category')
            revised=self.sync(conn,mappings={'Tyres':'Tyres'})
            self.assertEqual(revised['rows'][0]['operating_expenses'],Decimal('275.00'))

    def test_no_revenue_needs_explicit_no_activity_proof(self):
        with self.engine.begin() as conn:
            conn.execute(text('DELETE FROM dhl_orders'));self.raw=fixture.raw_book(conn);self.archive(conn)
            first=self.sync(conn,content=fixture.book(),filename='net.xlsx')
            self.assertIsNone(first['rows'][0]['gross'])
            second=self.sync(conn,no_activity={'KAA123B':'Synthetic dispatch log: no trips'})
            self.assertEqual(second['rows'][0]['commission'],Decimal('0.00'))
            self.assertEqual(second['rows'][0]['net_payable'],Decimal('-200.00'))

    def test_stale_upload_is_rejected_without_partial_save(self):
        with self.engine.begin() as conn:
            raw=self.archive(conn);first=self.sync(conn,content=fixture.book(),filename='net.xlsx')
            with self.assertRaises(VersionConflict):payout.synchronize(conn,fixture.MONTH,7,expected_report_id=None,expected_raw_id=raw['id'],content=fixture.book(200),filename='wrong.xlsx')
            self.assertEqual(payout.latest(conn,fixture.MONTH)['id'],first['id'])

    def test_reupload_of_older_version_warns_instead_of_claiming_replacement(self):
        with self.engine.begin() as conn:
            self.archive(conn);original=fixture.book()
            self.sync(conn,content=original,filename='net.xlsx')
            current=self.sync(conn,content=fixture.book(200),filename='corrected.xlsx')
            with self.assertRaisesRegex(VersionConflict,'earlier financial version'):
                self.sync(conn,content=original,filename='original-again.xlsx')
            self.assertEqual(payout.latest(conn,fixture.MONTH)['id'],current['id'])

    def test_source_dates_block_only_affected_truck(self):
        wb=load_workbook(io.BytesIO(self.raw));wb.active.append(['2025-02-10','OTHER999','OLD','Dist',50]);out=io.BytesIO();wb.save(out);wb.close()
        with self.engine.begin() as conn:
            self.archive(conn,out.getvalue());first=self.sync(conn,content=fixture.book(),filename='net.xlsx')
            self.assertIn('out_of_period_date',{i['code'] for i in first['issues']})
            confirm_truck(conn,first['rows'][0]['id'],7,'Synthetic this truck only')

    def test_revenue_without_expense_row_is_visible_not_fabricated_zero(self):
        with self.engine.begin() as conn:
            self.archive(conn)
            metadata,batch=payout.state(conn,fixture.MONTH)
            self.assertIsNone(batch);self.assertEqual(metadata['unmatched_trucks'][0]['gross'],Decimal('1000.00'))
            self.assertIsNone(metadata['unmatched_trucks'][0]['operating_expenses'])

    def test_month_totals_include_revenue_only_truck_without_fabricating_payout(self):
        with self.engine.begin() as conn:
            conn.execute(text("INSERT INTO dhl_orders(id,ref_no,date,truck_plate,distribution_cost,offloading_cost,total_revenue) VALUES(99,'RAWONLY','2026-02-10','RAWONLY999',200,0,200)"))
            self.raw=fixture.raw_book(conn);self.archive(conn)
            self.sync(conn,content=fixture.book(),filename='net.xlsx')
        status,body,_=self.request('GET','/dhl-reports/monthly-payouts/2026-02-01')
        self.assertEqual(status,200,body)
        self.assertEqual(body['report']['totals']['gross']['amount'],'1200.00')
        self.assertEqual(body['report']['totals']['net_payable']['missing_rows'],1)
        self.assertIsNone(body['report']['totals']['net_payable']['amount'])
        self.assertEqual(body['files']['unmatched_trucks'][0]['normalized_plate'],'RAWONLY999')

    def test_workbook_only_truck_is_visible_with_unverified_revenue(self):
        wb=load_workbook(io.BytesIO(self.raw));wb.active.append(['2026-02-10','WORKBOOK999','NOLEDGER','Dist',200]);out=io.BytesIO();wb.save(out);wb.close()
        with self.engine.begin() as conn:
            self.archive(conn,out.getvalue());self.sync(conn,content=fixture.book(),filename='net.xlsx')
            metadata,_=payout.state(conn,fixture.MONTH)
            row=next(r for r in metadata['unmatched_trucks'] if r['normalized_plate']=='WORKBOOK999')
            self.assertIsNone(row['gross']);self.assertIsNone(row['operating_expenses'])

    def test_facade_roles_and_sync_contract(self):
        with self.engine.begin() as conn:raw=self.archive(conn)
        for user,expected in [(None,401),(8,403),(9,403),(7,200)]:
            code,_,_=self.request('GET','/dhl-reports/monthly-payouts/2026-02-01',user=user)
            self.assertEqual(code,expected)
        code,body,_=self.request('POST','/dhl-reports/monthly-payouts/2026-02-01/expenses',content=fixture.book(),form_fields={'expected_raw_id':str(raw['id'])})
        self.assertEqual(code,200,body)
        code,again,_=self.request('POST','/dhl-reports/monthly-payouts/2026-02-01/sync',json_body={'expected_report_id':body['files']['latest_report_id'],'expected_raw_id':raw['id']})
        self.assertEqual(code,200,again);self.assertEqual(again['files']['latest_report_id'],body['files']['latest_report_id'])

    @unittest.skipUnless((Path(__file__).resolve().parents[1]/'data/JANUARY NET SUMMARY.xlsx').exists(), 'Actual January fixtures required')
    def test_january_month_view_reuses_reads_without_changing_any_results(self):
        from services.monthly_expense_dates import inspect_source, load_workbook
        from services.monthly_files import register
        from services.monthly_expense_parser import parse_expenses
        month=date(2026,1,1);root=Path(__file__).resolve().parents[1]/'data'
        raw=(root/'JANUARY 2026 TRIPS SUMMARY (1).xlsx').read_bytes();net=(root/'JANUARY NET SUMMARY.xlsx').read_bytes()
        _,manifest=inspect_source(raw,month)
        plates=sorted({k[1] for k in manifest}|{r.plate for r in parse_expenses(net,month).rows})
        with self.engine.begin() as conn:
            conn.execute(text('DELETE FROM dhl_orders'))
            for i,plate in enumerate(plates,100):
                conn.execute(text("INSERT INTO vehicles VALUES(:id,:plate,7,'30T')"),dict(id=i,plate=plate))
            for i,(key,amount) in enumerate(manifest.items(),1000):
                conn.execute(text('INSERT INTO dhl_orders VALUES(:id,:ref,:date,:plate,:vehicle,:dist,:off,:gross)'),
                    dict(id=i,ref=key[2],date=key[0],plate=key[1],vehicle=100+plates.index(key[1]),dist=amount[0],off=amount[1],gross=sum(amount)))
            saved=register(conn,raw,month,'january.xlsx',7)
            payout.synchronize(conn,month,7,expected_report_id=None,expected_raw_id=saved['id'],content=net,filename='net.xlsx')
            before=self.legacy(conn)
        url='/dhl-reports/monthly-payouts/2026-01-01'
        with patch('routers.monthly_payouts.reporting_read_scope',return_value=nullcontext()),patch('services.monthly_expense_dates.load_workbook',wraps=load_workbook) as opens:
            start=perf_counter();status,baseline,_=self.request('GET',url);old_seconds=perf_counter()-start;old_opens=opens.call_count
        self.assertEqual(status,200,baseline)
        with patch('services.monthly_expense_dates.load_workbook',wraps=load_workbook) as opens:
            start=perf_counter();status,optimized,_=self.request('GET',url);new_seconds=perf_counter()-start;new_opens=opens.call_count
        self.assertEqual(status,200,optimized);self.assertEqual(optimized,baseline)
        self.assertEqual(new_opens,2);self.assertGreater(old_opens,new_opens*10)
        with self.engine.begin() as conn:self.assertEqual(self.legacy(conn),before)
        # Local evidence only; never writes financial fixtures or production data.
        Path('/tmp/smithsons-month-loading-benchmark.json').write_text(json.dumps(dict(
            trucks=len(plates),source_rows=747,unoptimized_seconds=old_seconds,optimized_seconds=new_seconds,
            unoptimized_workbook_opens=old_opens,optimized_workbook_opens=new_opens,identical_response=True)))
