"""Disposable PostgreSQL, real immutable guards and scoped validation. No production."""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime,timedelta
from decimal import Decimal
import io
import unittest

from jose import jwt
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import text,select,update
from sqlalchemy.exc import DBAPIError

import test_monthly_expense_persistence as fixtures
import test_monthly_expense_api as api_fixtures
import test_monthly_statement_documents as document_fixtures
from services.monthly_expense_persistence import confirm_import,get_import,ValidationBlocked,VersionConflict
from services.monthly_expense_models import rows,imports
from services.monthly_expense_reconciliation import review_historical_adjustments
from services.truck_confirmations import confirm_truck,presentation,confirmations
from services.monthly_statement_documents import create_document,get_document


def two_trucks(raw=None):
    wb=load_workbook(io.BytesIO(raw or fixtures.book()))
    wb.active.append(['KBB 456C',25,0,0,0,0,0])
    stream=io.BytesIO();wb.save(stream);wb.close();return stream.getvalue()


@unittest.skipUnless(fixtures.URL,'Explicit disposable local PostgreSQL required')
class TruckTests(unittest.TestCase):
    setUpClass=classmethod(fixtures.PersistenceTests.setUpClass.__func__)
    tearDownClass=classmethod(fixtures.PersistenceTests.tearDownClass.__func__)
    tearDown=fixtures.PersistenceTests.tearDown
    draft=fixtures.PersistenceTests.draft
    legacy=fixtures.PersistenceTests.legacy

    def setUp(self):
        fixtures.PersistenceTests.setUp(self)
        with self.engine.begin() as conn:
            for name in ['002_monthly_statement_documents.sql','003_expense_zero_policy.sql','004_monthly_truck_confirmations.sql']:
                conn.exec_driver_sql((fixtures.ROOT/'migrations'/name).read_text())

    def test_good_truck_confirms_and_prints_despite_bad_other_truck(self):
        with self.engine.begin() as conn:
            before=self.legacy(conn)
            batch=self.draft(conn,content=two_trucks());row=batch['rows'][0]
            state=presentation(conn,batch)
            self.assertTrue(state['rows'][row['id']]['can_confirm'])
            self.assertFalse(state['rows'][batch['rows'][1]['id']]['can_confirm'])
            confirmation=confirm_truck(conn,row['id'],7,'Synthetic truck independently reviewed')
            self.assertEqual(confirm_truck(conn,row['id'],7,'Retry')['id'],confirmation['id'])
            document=create_document(conn,row['id'],7)
            content=PdfReader(io.BytesIO(bytes(document['pdf_content']))).pages[0].extract_text()
            self.assertIn('KES 744.00',content)
            self.assertIn('FUEL',content);self.assertIn('BREAKAGES',content)
            self.assertNotIn('Kasarani',content);self.assertNotIn('GPRS',content)
            self.assertIn('Total Expenses',content)
            self.assertEqual(self.legacy(conn),before)
            self.assertEqual(get_import(conn,batch['id'])['status'],'draft')
            with self.assertRaises(ValidationBlocked):confirm_import(conn,batch['id'],7)

    def test_confirmed_source_parent_and_all_its_rows_are_frozen(self):
        with self.engine.begin() as conn:
            batch=self.draft(conn,content=two_trucks());confirm_truck(conn,batch['rows'][0]['id'],7,'Reviewed')
        statements=['UPDATE dhl_monthly_expense_rows SET gross=12',
            "UPDATE dhl_monthly_expense_imports SET original_filename='changed.xlsx'",
            'DELETE FROM dhl_monthly_expense_rows','DELETE FROM dhl_monthly_expense_imports',
            'DELETE FROM dhl_monthly_truck_confirmations','TRUNCATE dhl_monthly_truck_confirmations',
            "UPDATE dhl_monthly_truck_confirmations SET reason='changed'",
            "UPDATE dhl_monthly_truck_confirmations SET status='superseded'"]
        for sql in statements:
            with self.subTest(sql=sql),self.assertRaises(DBAPIError),self.engine.begin() as conn:conn.execute(text(sql))

    def test_new_version_does_not_accumulate_or_revoke_other_confirmations(self):
        with self.engine.begin() as conn:
            first=self.draft(conn,content=two_trucks());a=first['rows'][0]
            original=confirm_truck(conn,a['id'],7,'A verified')
            old_pdf=create_document(conn,a['id'],7)
            # B receives a vehicle and explicit no-activity proof in a NEW version.
            conn.execute(text("INSERT INTO vehicles VALUES (2,'KBB456C',7,'30T')"))
            second=self.draft(conn,content=two_trucks(fixtures.book(200)),previous_version_id=first['id'],reason='Corrected expense workbook',
                verified_no_activity={'KBB456C':'Synthetic dispatch confirmed no business'})
            b=second['rows'][1];confirm_truck(conn,b['id'],7,'B verified independently')
            active=conn.execute(select(confirmations).where(confirmations.c.id==original['id'])).mappings().one()
            self.assertEqual(active['status'],'current')
            self.assertEqual(bytes(get_document(conn,old_pdf['id'])['pdf_content']),bytes(old_pdf['pdf_content']))
            with self.assertRaises(VersionConflict):confirm_truck(conn,second['rows'][0]['id'],7,'A correction')
            updated=confirm_truck(conn,second['rows'][0]['id'],7,'A correction explicitly approved',original['id'])
            self.assertEqual(updated['previous_confirmation_id'],original['id'])
            adjustment=updated['validation']['financial_adjustment']
            self.assertEqual(adjustment['net_payable_difference'],'-93.00')
            self.assertEqual(adjustment['expenses_difference'],'100.00')
            self.assertEqual(adjustment['commission_difference'],'-7.00')
            self.assertEqual(adjustment['gross_difference'],'0.00')
            self.assertFalse(adjustment['payment_processing'])
            self.assertEqual(adjustment['payment_status'],'not_assessed')
            self.assertNotIn('financial_adjustment',original['validation'])
            corrected_pdf=create_document(conn,second['rows'][0]['id'],7)
            self.assertIn('Separate correction adjustment: KES -93.00',PdfReader(io.BytesIO(bytes(corrected_pdf['pdf_content']))).pages[0].extract_text())

            self.assertEqual(second['rows'][0]['operating_expenses'],Decimal('300.00'))
            self.assertEqual(presentation(conn,second)['current_approved_totals']['operating_expenses'],Decimal('325.00'))
            self.assertEqual(create_document(conn,second['rows'][0]['id'],7)['previous_document_id'],old_pdf['id'])

    def test_unrelated_historical_dates_do_not_block_but_target_dates_do(self):
        with self.engine.begin() as conn:
            wb=load_workbook(io.BytesIO(fixtures.raw_book(conn)))
            wb.active.append(['2025-02-10','KBB456C','HIST','Dist',25])
            data=io.BytesIO();wb.save(data);wb.close()
            batch=self.draft(conn,historical_review=review_historical_adjustments(data.getvalue(),fixtures.MONTH))
            self.assertTrue(any(i['blocking'] for i in batch['issues']))
            confirm_truck(conn,batch['rows'][0]['id'],7,'Target dates verified; other truck pending')
        with self.engine.begin() as conn:
            first=get_import(conn,batch['id'])
            wb=load_workbook(io.BytesIO(fixtures.raw_book(conn)))
            wb.active.append(['2025-02-10','KAA123B','HIST','Dist',25])
            data=io.BytesIO();wb.save(data);wb.close()
            second=self.draft(conn,previous_version_id=first['id'],reason='Target historical review pending',
                historical_review=review_historical_adjustments(data.getvalue(),fixtures.MONTH))
            with self.assertRaises(ValidationBlocked):confirm_truck(conn,second['rows'][0]['id'],7,'Still pending',1)

    def test_unknown_truck_in_raw_date_issue_blocks_unsafe_scope(self):
        with self.engine.begin() as conn:
            wb=load_workbook(io.BytesIO(fixtures.raw_book(conn)));wb.active.append(['bad date',None,'UNKNOWN','Dist',25])
            data=io.BytesIO();wb.save(data);wb.close()
            batch=self.draft(conn,historical_review=review_historical_adjustments(data.getvalue(),fixtures.MONTH))
            with self.assertRaises(ValidationBlocked):confirm_truck(conn,batch['rows'][0]['id'],7,'Cannot attribute source issue')

    def test_other_truck_live_source_changes_do_not_block_unchanged_target(self):
        with self.engine.begin() as conn:
            conn.execute(text("INSERT INTO dhl_orders VALUES (2,'BREF','2026-02-10','KBB456C',NULL,500,0,500)"))
            batch=self.draft(conn)
            conn.execute(text('UPDATE dhl_orders SET distribution_cost=600,total_revenue=600 WHERE id=2'))
            confirm_truck(conn,batch['rows'][0]['id'],7,'A source unchanged, B source needs review')

    def test_whole_month_compatible_and_partial_correction_cannot_be_silently_replaced(self):
        with self.engine.begin() as conn:
            first=self.draft(conn);confirm_truck(conn,first['rows'][0]['id'],7,'Explicit A')
            confirm_import(conn,first['id'],7)
            second=self.draft(conn,content=fixtures.book(200),previous_version_id=first['id'],reason='Replacement')
            with self.assertRaises(VersionConflict):confirm_import(conn,second['id'],7)
            current=conn.execute(select(confirmations.c.id).where(confirmations.c.status=='current')).scalar_one()
            confirm_truck(conn,second['rows'][0]['id'],7,'Explicit correction',current)
            self.assertEqual(confirm_import(conn,second['id'],7)['status'],'confirmed')

    def test_existing_whole_month_correction_remains_supported(self):
        with self.engine.begin() as conn:
            first=self.draft(conn);confirm_import(conn,first['id'],7)
            second=self.draft(conn,content=fixtures.book(200),previous_version_id=first['id'],reason='Whole-month correction')
            confirm_import(conn,second['id'],7)
            self.assertEqual(conn.execute(select(confirmations.c.id).where(confirmations.c.status=='current')).all().__len__(),1)

    def test_concurrent_confirmation_is_idempotent(self):
        with self.engine.begin() as conn:batch=self.draft(conn)
        def confirm(_):
            with self.engine.begin() as conn:return confirm_truck(conn,batch['rows'][0]['id'],7,'Concurrent retry')['id']
        with ThreadPoolExecutor(max_workers=2) as pool:ids=list(pool.map(confirm,range(2)))
        self.assertEqual(ids[0],ids[1])

    def test_competing_versions_cannot_both_confirm_without_explicit_predecessor(self):
        with self.engine.begin() as conn:
            first=self.draft(conn)
            second=self.draft(conn,content=fixtures.book(200),previous_version_id=first['id'],reason='Competing candidate')
        def confirm(row_id):
            try:
                with self.engine.begin() as conn:return confirm_truck(conn,row_id,7,'Review candidate')['id']
            except VersionConflict:return None
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(confirm,[first['rows'][0]['id'],second['rows'][0]['id']]))
        self.assertEqual(sum(r is not None for r in results),1)

    def test_unicode_original_descriptions_and_item_sum_are_exact(self):
        with self.engine.begin() as conn:
            wb=load_workbook(io.BytesIO(fixtures.book(extra=True)));wb.active.cell(2,8,'Repairs – Côte')
            content=io.BytesIO();wb.save(content);wb.close()
            batch=self.draft(conn,content=content.getvalue(),approved_mappings={'Repairs – Côte':'Repairs'},reason='Category verified')
            row=batch['rows'][0];confirm_truck(conn,row['id'],7,'All expenses verified')
            doc=create_document(conn,row['id'],7)
            content=PdfReader(io.BytesIO(bytes(doc['pdf_content']))).pages[0].extract_text()
            self.assertIn('Repairs – Côte',content)
            self.assertNotIn('GPRS',content)
            import re
            segment=content.split('Gross Transport Revenue')[1].split('Total Expenses')[0]
            amounts=[Decimal(a.replace(',','')) for a in re.findall(r'KES ([\d,.-]+)',segment)][1:]
            self.assertEqual(sum(amounts),row['operating_expenses'])

    def test_zero_activity_negative_balance_and_exact_nonzero_dynamic_items(self):
        with self.engine.begin() as conn:
            conn.execute(text('DELETE FROM dhl_orders'))
            batch=self.draft(conn,content=fixtures.book(blank=True,extra=True),approved_mappings={'Tyres':'Tyres'},reason='Synthetic category approved',
                verified_no_activity={'KAA123B':'Dispatch reviewed'})
            row=batch['rows'][0];confirm_truck(conn,row['id'],7,'Zero business and all expenses verified')
            pdf=create_document(conn,row['id'],7)
            content=PdfReader(io.BytesIO(bytes(pdf['pdf_content']))).pages[0].extract_text()
            self.assertEqual(row['commission'],Decimal('0.00'));self.assertEqual(row['net_payable'],Decimal('-175.00'))
            self.assertIn('Tyres',content);self.assertNotIn('FUEL',content)
            self.assertEqual(sum(Decimal(v['amount']) for v in row['deductions'].values() if Decimal(v['amount'])!=0),row['operating_expenses'])

    def test_replacement_omitting_a_truck_keeps_its_current_approval_visible(self):
        with self.engine.begin() as conn:
            first=self.draft(conn);confirm_truck(conn,first['rows'][0]['id'],7,'A verified')
            conn.execute(text("INSERT INTO vehicles VALUES (2,'KBB456C',7,'30T')"))
            second=self.draft(conn,content=fixtures.book(fuel=0,plate='KBB456C'),previous_version_id=first['id'],reason='Replacement roster omits A',
                verified_no_activity={'KBB456C':'Dispatch verified'})
            state=presentation(conn,second)
            self.assertEqual(state['current_approvals'][0]['normalized_plate'],'KAA123B')
            self.assertEqual(state['current_approved_totals']['net_payable'],Decimal('744.00'))
            self.assertEqual(second['rows'][0]['operating_expenses'],Decimal('100.00'))

    def test_hashes_and_pdf_verification_survive_connection_timezone_changes(self):
        with self.engine.begin() as conn:
            batch=self.draft(conn);confirm_truck(conn,batch['rows'][0]['id'],7,'Verified')
        with self.engine.begin() as conn:
            conn.execute(text("SET LOCAL TIME ZONE 'Africa/Nairobi'"))
            self.assertTrue(create_document(conn,batch['rows'][0]['id'],7)['pdf_sha256'])


@unittest.skipUnless(fixtures.URL,'Explicit disposable local PostgreSQL required')
class TruckApiTests(unittest.TestCase):
    setUpClass=classmethod(fixtures.PersistenceTests.setUpClass.__func__)
    tearDownClass=classmethod(fixtures.PersistenceTests.tearDownClass.__func__)
    tearDown=fixtures.PersistenceTests.tearDown
    request=api_fixtures.ApiTests.request
    save=api_fixtures.ApiTests.save

    def setUp(self):
        document_fixtures.StatementApiTests.setUp(self)
        with self.engine.begin() as conn:
            for name in ['003_expense_zero_policy.sql','004_monthly_truck_confirmations.sql']:
                conn.exec_driver_sql((fixtures.ROOT/'migrations'/name).read_text())

    def post(self,path,role=7,**payload):
        import asyncio
        import json
        import contextlib
        async def run():
            headers=[]
            if role is not None:
                token=jwt.encode({'sub':str(role),'exp':datetime.utcnow()+timedelta(minutes=5)},'isolated-phase3-test-only',algorithm='HS256')
                headers=[(b'authorization',('Bearer '+token).encode())]
            body=json.dumps(payload).encode();headers.extend([(b'content-type',b'application/json'),(b'content-length',str(len(body)).encode())])
            messages=[]
            async def receive():return {'type':'http.request','body':body,'more_body':False}
            async def send(message):messages.append(message)
            await self.app({'type':'http','asgi':{'version':'3.0'},'http_version':'1.1','method':'POST','scheme':'http',
                'path':path,'raw_path':path.encode(),'query_string':b'','headers':headers,'client':('127.0.0.1',1),'server':('test',80)},receive,send)
            from types import SimpleNamespace
            return SimpleNamespace(status_code=messages[0]['status'],text=b''.join(m.get('body',b'') for m in messages).decode())
        return asyncio.run(run())

    def test_admin_only_confirmation_and_report_states(self):
        batch=self.save(content=two_trucks());row=batch['rows'][0]
        self.assertTrue(row['truck_confirmation']['can_confirm'])
        path=f"/dhl-reports/expense-rows/{row['id']}/confirm"
        for role,expected in [(None,401),(8,403),(9,403)]:self.assertEqual(self.post(path,role=role,reason='Reviewed').status_code,expected)
        self.assertEqual(self.post(path,reason='').status_code,422)
        self.assertEqual(self.post(path,reason='Reviewed',gross='999').status_code,422)
        confirmed=self.post(path,reason='Admin independently verified truck');self.assertEqual(confirmed.status_code,200,confirmed.text)
        detail=self.request('GET',f"/dhl-reports/expense-imports/{batch['id']}")[1]
        self.assertEqual(detail['status'],'draft');self.assertEqual(detail['truck_confirmation_summary']['confirmed_count'],1)
        self.assertEqual(detail['rows'][0]['truck_confirmation']['status'],'confirmed')
        self.assertEqual(self.request('POST',f"/dhl-reports/monthly-statement-documents/from-row/{row['id']}")[0],200)


@unittest.skipUnless(fixtures.URL,'Explicit disposable local PostgreSQL required')
class MigrationTests(unittest.TestCase):
    setUpClass=classmethod(fixtures.PersistenceTests.setUpClass.__func__)
    tearDownClass=classmethod(fixtures.PersistenceTests.tearDownClass.__func__)
    tearDown=fixtures.PersistenceTests.tearDown
    draft=fixtures.PersistenceTests.draft

    def setUp(self):
        fixtures.PersistenceTests.setUp(self)
        with self.engine.begin() as conn:
            for name in ['002_monthly_statement_documents.sql','003_expense_zero_policy.sql']:
                conn.exec_driver_sql((fixtures.ROOT/'migrations'/name).read_text())

    def test_backfill_preserves_existing_financial_rows_and_pdf_bytes(self):
        with self.engine.begin() as conn:
            batch=self.draft(conn);old=confirm_import(conn,batch['id'],7)
            pdf=create_document(conn,old['rows'][0]['id'],7)
        with self.engine.begin() as conn:
            conn.exec_driver_sql((fixtures.ROOT/'migrations/004_monthly_truck_confirmations.sql').read_text())
            self.assertEqual(get_import(conn,batch['id']),old)
            self.assertEqual(get_document(conn,pdf['id'])['pdf_content'],pdf['pdf_content'])
            self.assertEqual(conn.execute(select(confirmations.c.method)).scalar_one(),'month')

    def test_migration_is_atomic_including_function_replacements(self):
        with self.engine.begin() as conn:
            original=conn.execute(text("SELECT pg_get_functiondef('monthly_statement_guard()'::regprocedure)")).scalar_one()
        with self.assertRaises(DBAPIError),self.engine.begin() as conn:
            conn.exec_driver_sql((fixtures.ROOT/'migrations/004_monthly_truck_confirmations.sql').read_text())
            conn.execute(text('SELECT 1/0'))
        with self.engine.begin() as conn:
            self.assertIsNone(conn.execute(text("SELECT to_regclass('dhl_monthly_truck_confirmations')")).scalar_one())
            self.assertEqual(conn.execute(text("SELECT pg_get_functiondef('monthly_statement_guard()'::regprocedure)")).scalar_one(),original)
