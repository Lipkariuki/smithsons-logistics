"""Saved files and source reuse in isolated PostgreSQL; no production imports."""
from concurrent.futures import ThreadPoolExecutor
from datetime import date
import io
import unittest
from openpyxl import load_workbook
from sqlalchemy import text,select
from sqlalchemy.exc import DBAPIError
import test_monthly_expense_api as api
import test_monthly_expense_persistence as fixtures
from services import monthly_files as files
from services.monthly_expense_persistence import create_draft,confirm_import,get_import,VersionConflict
from services.monthly_statement_documents import create_document,get_document
from services.truck_confirmations import confirm_truck
from routers.monthly_files import create_router
from routers.monthly_statement_documents import create_router as statements_router


@unittest.skipUnless(fixtures.URL,'Explicit disposable local PostgreSQL required')
class SavedFileTests(unittest.TestCase):
 setUpClass=classmethod(api.ApiTests.setUpClass.__func__)
 tearDownClass=classmethod(api.ApiTests.tearDownClass.__func__)
 tearDown=api.ApiTests.tearDown
 request=api.ApiTests.request
 legacy=api.ApiTests.legacy
 def setUp(self):
  api.ApiTests.setUp(self)
  with self.engine.begin() as conn:
   for name in ['002_monthly_statement_documents.sql','003_expense_zero_policy.sql','004_monthly_truck_confirmations.sql','005_saved_dhl_workbooks.sql']:
    conn.exec_driver_sql((fixtures.ROOT/'migrations'/name).read_text())
   self.raw=fixtures.raw_book(conn)
  def dependency():
   with self.engine.begin() as conn:yield conn
  self.app.include_router(create_router(dependency,api.auth_functions(self.engine)))
  self.app.include_router(statements_router(dependency,api.auth_functions(self.engine)))
 def archive(self,conn,content=None,**kwargs):return files.register(conn,content or self.raw,fixtures.MONTH,'original.xlsx',7,**kwargs)
 def test_retention_deduplication_and_immutable_metadata(self):
  with self.engine.begin() as conn:
   before=self.legacy(conn);saved=self.archive(conn)
   self.assertEqual(self.archive(conn)['id'],saved['id']);self.assertEqual(saved['content'],self.raw)
   self.assertEqual(files.month_files(conn,fixtures.MONTH)['raw_files'][0]['status'],'uploaded')
   self.assertEqual(self.legacy(conn),before)
  for command in ['UPDATE dhl_monthly_workbooks SET reason=\'changed\'','DELETE FROM dhl_monthly_workbooks','TRUNCATE dhl_monthly_workbooks']:
   with self.subTest(command=command),self.assertRaises(DBAPIError),self.engine.begin() as conn:conn.execute(text(command))
 def test_replacements_require_latest_id_and_reason(self):
  book=load_workbook(io.BytesIO(self.raw));book.active['E2']=901;stream=io.BytesIO();book.save(stream);book.close()
  with self.engine.begin() as conn:
   first=self.archive(conn)
   with self.assertRaises(VersionConflict):self.archive(conn,stream.getvalue())
   with self.assertRaises(ValueError):self.archive(conn,stream.getvalue(),previous_id=first['id'])
   second=self.archive(conn,stream.getvalue(),previous_id=first['id'],reason='Correct source')
   self.assertEqual(second['version'],2);self.assertEqual(second['previous_id'],first['id'])
   self.assertEqual(files.get(conn,first['id'])['content'],self.raw)
   self.assertEqual(files.month_files(conn,fixtures.MONTH)['raw_files'][0]['status'],'needs_review')
 def test_incomplete_source_is_not_missing_and_requires_explicit_review(self):
  with self.engine.begin() as conn:
   original=create_draft(conn,fixtures.book(),fixtures.MONTH,'net.xlsx',7)
   self.assertEqual(original['rows'][0]['revenue_status'],'incomplete');self.assertIsNone(original['rows'][0]['gross'])
   raw=self.archive(conn)
   # Saving a source does not change the original report or approve it.
   self.assertIsNone(get_import(conn,original['id'])['rows'][0]['gross'])
   meta=files.month_files(conn,fixtures.MONTH);self.assertEqual(meta['parsed_dhl']['gross'],'1000.00');self.assertFalse(meta['source_verified'])
   revised=create_draft(conn,fixtures.book(),fixtures.MONTH,'net.xlsx',7,previous_version_id=original['id'],reason='Source reviewed',source_verification='Synthetic dispatch and source checked',historical_review=files.reference(raw))
   self.assertEqual(str(revised['rows'][0]['gross']),'1000.00');self.assertNotIn('source_content_base64',revised['decisions']['historical_review'])
   confirmed=confirm_truck(conn,revised['rows'][0]['id'],7,'Synthetic truck verified')
   document=create_document(conn,revised['rows'][0]['id'],7)
   self.assertEqual(bytes(get_document(conn,document['id'])['pdf_content']),bytes(document['pdf_content']))
   self.assertIsNone(get_import(conn,original['id'])['rows'][0]['gross'])
 def test_month_mismatch_tampering_and_historical_dates_block(self):
  with self.engine.begin() as conn:
   archive=self.archive(conn)
   with self.assertRaises(ValueError):files.resolve_review(conn,files.reference(archive),date(2026,3,1))
   with self.assertRaises(ValueError):files.resolve_review(conn,dict(files.reference(archive),sha256='0'*64),fixtures.MONTH)
   wb=load_workbook(io.BytesIO(self.raw));wb.active.append(['2025-02-10','KAA123B','HIST','Dist',20]);stream=io.BytesIO();wb.save(stream);wb.close()
   nextfile=self.archive(conn,stream.getvalue(),previous_id=archive['id'],reason='Historical source retained without approval')
   draft=create_draft(conn,fixtures.book(),fixtures.MONTH,'net.xlsx',7,source_verification='Synthetic source checked',historical_review=files.reference(nextfile))
   self.assertIn('out_of_period_date',{i['code'] for i in draft['issues']})
   from services.monthly_expense_persistence import ValidationBlocked
   with self.assertRaises(ValidationBlocked):confirm_truck(conn,draft['rows'][0]['id'],7,'Cannot approve historical source')
 def test_admin_only_metadata_and_workbook_downloads(self):
  with self.engine.begin() as conn:
   raw=self.archive(conn);draft=create_draft(conn,fixtures.book(),fixtures.MONTH,'net.xlsx',7,historical_review=files.reference(raw))
  for path in ['/dhl-reports/monthly-files?month=2026-02-01',f"/dhl-reports/monthly-files/raw/{raw['id']}",f"/dhl-reports/monthly-files/raw/{raw['id']}/verification",f"/dhl-reports/monthly-files/report/{draft['id']}/expenses",f"/dhl-reports/monthly-files/report/{draft['id']}/raw"]:
   for user,expected in [(None,401),(8,403),(9,403),(7,200)]:
    status,body,headers=self.request('GET',path,user=user);self.assertEqual(status,expected,(path,user,body))
    if user==7:self.assertEqual(headers.get('cache-control'),'no-store')
  for user,expected in [(None,401),(8,403),(9,403),(7,200)]:
   code,body,_=self.request('POST','/dhl-reports/monthly-files/dhl',content=self.raw,user=user,form_fields={})
   self.assertEqual(code,expected,body)
  status,content,headers=self.request('GET',f"/dhl-reports/monthly-files/report/{draft['id']}/raw")
  self.assertEqual(content,self.raw);self.assertIn('attachment',headers['content-disposition'])
 def test_preview_and_saved_revision_reuse_archived_raw_and_original_net(self):
  with self.engine.begin() as conn:
   raw=self.archive(conn);draft=create_draft(conn,fixtures.book(),fixtures.MONTH,'net.xlsx',7)
  # Existing saved expense bytes are reused, no uploaded file or raw file supplied.
  fields={'mode':'save','expected_file_sha256':draft['file_sha256']}
  code,result,_=self.request('POST',f"/dhl-reports/expense-imports/{draft['id']}/revise",options={'reason':'Review saved files','source_verification':'Synthetic source checked','raw_workbook_id':raw['id']},form_fields=fields)
  self.assertEqual(code,200,result);self.assertEqual(result['version'],2);self.assertEqual(result['rows'][0]['gross'],'1000.00');self.assertTrue(result['can_confirm'],result['issues'])
  self.assertEqual(result['decisions']['historical_review']['archive_id'],raw['id']);self.assertEqual(result['raw_source_file']['version'],1)
 def test_revised_source_cannot_inherit_old_historical_exclusions(self):
  with self.engine.begin() as conn:
   wb=load_workbook(io.BytesIO(self.raw));wb.active.append(['2025-02-10','KAA123B','OLD-HISTORY','Dist',20]);out=io.BytesIO();wb.save(out)
   first=self.archive(conn,out.getvalue());choices={'4':{'action':'exclude_from_period','reason':'Synthetic original source reviewed'}}
   original=create_draft(conn,fixtures.book(),fixtures.MONTH,'net.xlsx',7,source_verification='Reviewed',historical_review=files.reference(first),date_dispositions=choices)
   wb.active['C4']='DIFFERENT-HISTORY';out=io.BytesIO();wb.save(out);wb.close()
   second=self.archive(conn,out.getvalue(),previous_id=first['id'],reason='Different source')
  status,body,_=self.request('POST',f"/dhl-reports/expense-imports/{original['id']}/revise",options={'reason':'Different raw source','raw_workbook_id':second['id'],'date_dispositions':choices},form_fields={'mode':'preview','expected_file_sha256':original['file_sha256']})
  self.assertEqual(status,422,body)

 def test_atomic_migration_rollback_preserves_schema(self):
  with self.assertRaises(DBAPIError),self.engine.begin() as conn:
   conn.exec_driver_sql('CREATE TABLE transient_file_test(id int); SELECT 1/0;')
  with self.engine.begin() as conn:self.assertIsNone(conn.execute(text("SELECT to_regclass('transient_file_test')")).scalar())
 def test_concurrent_same_file_uploads_are_idempotent(self):
  def save(_):
   with self.engine.begin() as conn:return self.archive(conn)['id']
  with ThreadPoolExecutor(max_workers=2) as pool:ids=list(pool.map(save,range(2)))
  self.assertEqual(ids[0],ids[1])

 def test_attach_matching_evidence_preserves_manual_payslips_and_reuses_net(self):
  with self.engine.begin() as conn:
   before=self.legacy(conn);draft=create_draft(conn,fixtures.book(),fixtures.MONTH,'saved-net.xlsx',7)
  status,saved,_=self.request('POST','/dhl-reports/monthly-files/dhl',content=self.raw)
  self.assertEqual(status,200,saved);self.assertTrue(saved['verification']['matches'])
  self.assertEqual(saved['verification']['saved_totals']['gross'],'1000.00')
  status,check,_=self.request('GET',f"/dhl-reports/monthly-files/raw/{saved['id']}/verification")
  self.assertEqual(status,200);self.assertTrue(check['matches'])
  with self.engine.begin() as conn:
   self.assertEqual(self.legacy(conn),before)
   self.assertIsNone(get_import(conn,draft['id'])['rows'][0]['gross'])
   self.assertFalse(files.month_files(conn,fixtures.MONTH)['source_verified'])
  # Explicit admin source attestation and NEW draft; neither workbook is uploaded again.
  status,revised,_=self.request('POST',f"/dhl-reports/expense-imports/{draft['id']}/revise",options={
   'source_verification':'Synthetic source and dispatch checked','reason':'Attached original evidence reviewed',
   'raw_workbook_id':saved['id']},form_fields={'mode':'save','expected_file_sha256':draft['file_sha256']})
  self.assertEqual(status,200,revised);row=revised['rows'][0]
  self.assertEqual([row[k] for k in ['gross','operating_expenses','commission','net_payable']],['1000.00','200.00','56.00','744.00'])
  self.assertEqual(revised['status'],'draft');self.assertIsNone(row['truck_confirmation']['record'])
  with self.engine.begin() as conn:self.assertEqual(self.legacy(conn),before)

 def test_evidence_equal_gross_cannot_hide_distribution_or_identity_mismatches(self):
  for cell,value,expected_reason in [('E2',899,'transport_amount_mismatch'),('C2','WRONG-REF','missing_saved_record'),('B2','OTHER999','missing_saved_record')]:
   with self.subTest(cell=cell):
    wb=load_workbook(io.BytesIO(self.raw));wb.active[cell]=value
    if cell=='E2':wb.active['E3']=101 # Same gross, wrong transport components.
    out=io.BytesIO();wb.save(out);wb.close()
    with self.engine.begin() as conn:
     # New evidence version per subtest, without replacing a single saved DHL row.
     latest=conn.execute(select(files.workbooks.c.id).order_by(files.workbooks.c.version.desc()).limit(1)).scalar_one_or_none()
     before=self.legacy(conn);archive=self.archive(conn,out.getvalue(),previous_id=latest,reason='Synthetic mismatch audit')
     check=files.verification(conn,archive);self.assertFalse(check['matches']);self.assertIn(expected_reason,{r['reason'] for r in check['differences']})
     draft=create_draft(conn,fixtures.book(fuel=101+len(check['differences'])),fixtures.MONTH,'net.xlsx',7,
      source_verification='Claim cannot bypass evidence checks',historical_review=files.reference(archive),
      previous_version_id=conn.execute(text('SELECT id FROM dhl_monthly_expense_imports ORDER BY version DESC LIMIT 1')).scalar_one_or_none(),reason='Mismatch test')
     self.assertIsNone(draft['rows'][0]['gross'])
     from services.monthly_expense_persistence import ValidationBlocked
     with self.assertRaises(ValidationBlocked):confirm_truck(conn,draft['rows'][0]['id'],7,'Must remain blocked')
     self.assertEqual(self.legacy(conn),before)

 def test_strict_year_and_month_review_is_not_approved_by_attachment(self):
  wb=load_workbook(io.BytesIO(self.raw));wb.active['A2']='2025-02-10';out=io.BytesIO();wb.save(out);wb.close()
  with self.engine.begin() as conn:
   archive=self.archive(conn,out.getvalue());check=files.verification(conn,archive)
   self.assertFalse(check['matches']);self.assertEqual(check['date_issues'][0]['source_row'],2)
   self.assertEqual(check['date_issues'][0]['code'],'out_of_period_date')
   self.assertEqual(check['date_issues'][0]['reference'],'REF1')
   self.assertEqual(check['workbook_totals']['gross'],'100.00');self.assertEqual(check['saved_totals']['gross'],'1000.00')

 def test_admin_audit_entry_does_not_reduce_verified_gross(self):
  wb=load_workbook(io.BytesIO(self.raw));wb.active.append(['2026-02-10','KAA 123B','REF1','Admin',-250]);out=io.BytesIO();wb.save(out);wb.close()
  with self.engine.begin() as conn:
   before=self.legacy(conn);archive=self.archive(conn,out.getvalue());check=files.verification(conn,archive)
   self.assertTrue(check['matches']);self.assertEqual(check['workbook_totals']['gross'],'1000.00');self.assertEqual(self.legacy(conn),before)

 def test_other_truck_date_discrepancy_does_not_block_verified_truck(self):
  wb=load_workbook(io.BytesIO(self.raw));wb.active.append(['2025-02-10','OTHER999','UNRESOLVED','Dist',75]);out=io.BytesIO();wb.save(out);wb.close()
  with self.engine.begin() as conn:
   archive=self.archive(conn,out.getvalue());self.assertFalse(files.verification(conn,archive)['matches'])
   draft=create_draft(conn,fixtures.book(),fixtures.MONTH,'net.xlsx',7,source_verification='Synthetic KAA123B source checked',historical_review=files.reference(archive))
   self.assertEqual(str(draft['rows'][0]['gross']),'1000.00')
   self.assertIn('out_of_period_date',{i['code'] for i in draft['issues']})
   confirmed=confirm_truck(conn,draft['rows'][0]['id'],7,'Only this eligible truck reviewed')
   self.assertEqual(confirmed['normalized_plate'],'KAA123B')

@unittest.skipUnless(fixtures.URL,'Explicit disposable local PostgreSQL required')
class SavedFileMigrationTests(unittest.TestCase):
 setUpClass=classmethod(fixtures.PersistenceTests.setUpClass.__func__)
 tearDownClass=classmethod(fixtures.PersistenceTests.tearDownClass.__func__)
 setUp=fixtures.PersistenceTests.setUp
 tearDown=fixtures.PersistenceTests.tearDown
 def test_exact_migration_rollback_and_one_shot_replay(self):
  migration=(fixtures.ROOT/'migrations/005_saved_dhl_workbooks.sql').read_text()
  with self.engine.begin() as conn:before=fixtures.PersistenceTests.legacy(self,conn)
  with self.assertRaises(DBAPIError),self.engine.begin() as conn:
   conn.exec_driver_sql(migration);conn.exec_driver_sql('SELECT 1/0')
  with self.engine.begin() as conn:
   self.assertIsNone(conn.execute(text("SELECT to_regclass('dhl_monthly_workbooks')")).scalar())
   self.assertEqual(fixtures.PersistenceTests.legacy(self,conn),before)
   conn.exec_driver_sql(migration)
  with self.assertRaises(DBAPIError),self.engine.begin() as conn:conn.exec_driver_sql(migration)
  with self.engine.begin() as conn:
   self.assertEqual(fixtures.PersistenceTests.legacy(self,conn),before)
   self.assertEqual(conn.execute(text('SELECT count(*) FROM dhl_monthly_workbooks')).scalar_one(),0)
