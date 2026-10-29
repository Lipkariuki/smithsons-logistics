"""Read-only workbook audit and actual legacy functions against in-memory SQLite.
Never import the production database/auth modules; no external network or SQL target.
The month-only in-memory excerpt is for testing, never submitted to an upload endpoint.
"""
from pathlib import Path
import sys, types, io, json, csv, socket
from collections import defaultdict
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
def deny_network(*args,**kwargs):raise RuntimeError('Revenue verification prohibits network access')
socket.socket.connect=deny_network
from sqlalchemy import create_engine,text,Date
from sqlalchemy.orm import declarative_base,sessionmaker
safe=types.ModuleType('database');safe.Base=declarative_base();safe.engine=create_engine('sqlite:///:memory:');safe.SessionLocal=sessionmaker(bind=safe.engine)
def get_db():
 with safe.SessionLocal() as session:yield session
safe.get_db=get_db;sys.modules['database']=safe
# Dependencies of real report route decorators only; actual auth is not under test.
auth=types.ModuleType('routers.auth');auth.require_role=lambda *roles: lambda:None;sys.modules['routers.auth']=auth
from models import User,Vehicle,DHLOrder
from services.dhl_import import parse_dhl_excel
from routers import dhl_reports
from services.monthly_expense_parser import parse_expenses,money
from services.monthly_expense_persistence import _sources
from services.monthly_expense_calculations import RevenueEvidence,RevenueStatus,calculate_monthly_expenses
from openpyxl import load_workbook,Workbook
MONTH=date(2026,2,1);END=date(2026,2,28);q=lambda n:n.quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
source=ROOT/'data/FEBRUARY 2026 TRIPS SUMMARY.xlsx';wb=load_workbook(source,data_only=True);ws=wb.active
headers={str(c.value).strip():c.column-1 for c in ws[1] if c.value is not None}
independent=defaultdict(lambda:dict(dist=Decimal(0),offloading=Decimal(0),admin=Decimal(0),admin_rows=0))
historical=[];admin=[];monthrows=[]
for n,r in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
 day=r[headers['DATE']]
 if not day:continue
 plate=str(r[headers['TRUCK']]).replace(' ','').upper();label=str(r[headers['COST']]).strip().lower();amount=Decimal(str(r[headers['AMOUNT']]))
 if (day.year,day.month)!=(2026,2):
  historical.append(dict(excel_row=n,date=day.isoformat(),truck=plate,reference=str(r[headers['REF NO']]),cost=r[headers['COST']],amount=str(amount)));continue
 monthrows.append(r)
 # Deliberately exact recognized labels for independent comparison, not importer substring logic.
 item=independent[plate]
 if label=='dist':item['dist']+=amount
 elif label=='offloading':item['offloading']+=amount
 elif label=='admin':
  item['admin']+=amount;item['admin_rows']+=1;admin.append(dict(excel_row=n,date=day.isoformat(),truck=plate,reference=str(r[headers['REF NO']]),amount=str(amount)))
 else:raise AssertionError('Unexpected February COST '+label)
wb.close()
# Full original workbook's unchanged validation must still reject mixed months.
try:parse_dhl_excel(source.read_bytes());raise AssertionError('Expected mixed-month rejection')
except ValueError as e:full_import_error=str(e)
# Preserve every February row including Admin in a clearly test-only in-memory excerpt.
excerpt=Workbook();sheet=excerpt.active
sheet.append([str(c) for c in headers])
for r in monthrows:sheet.append(list(r[:len(headers)]))
buffer=io.BytesIO();excerpt.save(buffer);excerpt.close()
parsed,start,end=parse_dhl_excel(buffer.getvalue());assert (start,end)==(MONTH,END)
safe.Base.metadata.create_all(safe.engine)
with safe.SessionLocal() as db:
 db.add(User(id=1,name='Read-only test owner',password_hash='unused',role='owner'));db.flush()
 vehicles={p:Vehicle(plate_number=p,owner_id=1) for p in independent};db.add_all(vehicles.values());db.flush()
 for row in parsed:db.add(DHLOrder(**row,vehicle_id=vehicles[row['truck_plate']].id))
 db.commit()
 args=dict(owner_id=None,vehicle_id=None,vehicle_plate=None,start_date=MONTH,end_date=END,db=db,current_user=None)
 legacy={r.plate_number:r for r in dhl_reports.dhl_summary(**args)}
 reconciliation={r.plate_number:r for r in dhl_reports.dhl_reconciliation(**args).rows}
 payslip_gross={p:dhl_reports._sum_dhl_revenue(db,v.id,MONTH,END) for p,v in vehicles.items()}
# Actual monthly-source query against local tables; SQLite lacks PostgreSQL SHARE lock.
# Skip only that lock statement; real SELECT, bounds, snapshot formatting and hashing run.
class ReadConnection:
 def __init__(self,conn):self.conn=conn
 def execute(self,statement,*args,**kwargs):
  if str(statement)=='LOCK TABLE dhl_orders IN SHARE MODE':return None
  if str(statement).startswith('SELECT id, ref_no, date, truck_plate'):statement=statement.columns(date=Date())
  return self.conn.execute(statement,*args,**kwargs)
with safe.engine.connect() as conn:records,fingerprint=_sources(ReadConnection(conn),MONTH)
monthly=defaultdict(lambda:Decimal(0));recordrounddelta=Decimal(0)
for r in records:
 monthly[r['truck_plate']]+=money(r['total_revenue'])
 recordrounddelta+=money(r['total_revenue'])-Decimal(r['total_revenue'])
imported=defaultdict(lambda:dict(dist=Decimal(0),offloading=Decimal(0),gross=Decimal(0)))
for r in parsed:
 for out,key in [('dist','distribution_cost'),('offloading','offloading_cost'),('gross','total_revenue')]:imported[r['truck_plate']][out]+=Decimal(str(r[key]))
rows=[]
for p,v in sorted(independent.items()):
 gross=q(v['dist']+v['offloading']);r=legacy[p]
 item=dict(truck=p,dist=q(v['dist']),offloading=q(v['offloading']),gross=gross,admin_audit_only=q(v['admin']),admin_rows=v['admin_rows'],importer_dist=q(imported[p]['dist']),importer_offloading=q(imported[p]['offloading']),importer_gross=q(imported[p]['gross']),report_gross=q(Decimal(str(r.total_revenue))),reconciliation_gross=q(Decimal(str(reconciliation[p].dhl_revenue))),payslip_gross=q(Decimal(str(payslip_gross[p]))),monthly_expense_gross=q(monthly[p]))
 for key in ['importer_dist','importer_offloading']:item[key+'_difference']=item[key]-item[key.removeprefix('importer_')]
 for key in ['importer_gross','report_gross','reconciliation_gross','payslip_gross','monthly_expense_gross']:item[key+'_difference']=item[key]-gross
 item['raw_float_report_difference']=str(Decimal(str(r.total_revenue))-gross)
 rows.append(item)
# Focused isolated fixtures prove charge handling with conflicting RATE/QTY and negative Admin.
def fixture(labels):
 w=Workbook();s=w.active;s.append(['DATE','REF NO','TRUCK','COST','AMOUNT','RATE','QTY'])
 for label,amount in labels:s.append([MONTH,'SAME-REF','SYN 123A',label,amount,999999,999999])
 b=io.BytesIO();w.save(b);w.close();return parse_dhl_excel(b.getvalue())[0]
synthetic=fixture([('Dist',100),('DIST',50),('Offloading',20),('OFFLOADING',5),('Admin',-9999)])
assert len(synthetic)==1
synthetic_admin_only=fixture([('Admin',-9999)])
# Unknown COST-like columns in a Net Summary cannot automatically become deductions.
w=Workbook();s=w.active;s.append([None,'FEBRUARY 2026 NET SUMMARY']);s.append(['TRUCK REG/NO','FUEL','DIST','Offloading','Admin']);s.append(['SYN123A',100,999,888,-777]);b=io.BytesIO();w.save(b);w.close();unknown=parse_expenses(b.getvalue(),MONTH)
ignored=parse_expenses(b.getvalue(),MONTH,approved_mappings={'DIST':None,'Offloading':None,'Admin':None})
calc=calculate_monthly_expenses(ignored.rows[0],MONTH,RevenueEvidence('SYN123A',MONTH,RevenueStatus.DHL_RECORDS,Decimal('1000.00'),'Synthetic fixture only'))
result=dict(full_original_import_error=full_import_error,month_source_rows=len(monthrows),historical=historical,admin_rows=admin,truck_count=len(rows),totals={key:str(sum((r[key] for r in rows),Decimal(0))) for key in ['dist','offloading','gross','admin_audit_only','importer_gross','report_gross','reconciliation_gross','payslip_gross','monthly_expense_gross']},nonzero_cent_differences=[{k:str(v) for k,v in r.items()} for r in rows if any(r[k]!=0 for k in r if k.endswith('_difference') and k!='raw_float_report_difference')],synthetic=synthetic,synthetic_admin_only=synthetic_admin_only,unknown_headers={h.original_header:h.classification for h in unknown.headers},unknown_deductions={k:str(v) for k,v in unknown.rows[0].expenses.items()},ignored_deductions={k:str(v) for k,v in ignored.rows[0].expenses.items()},synthetic_calculation={k:str(getattr(calc,k)) for k in ['gross','operating_expenses','commission_base','commission','net_payable']},source_fingerprint=fingerprint,rows=rows)
if '--json' in sys.argv:print(json.dumps(result,default=str));sys.exit()
out=ROOT/'output/dhl-revenue-verification';out.mkdir(exist_ok=True,parents=True)
(out/'verification.json').write_text(json.dumps(result,indent=2,default=str))
for filename,values in [('february-truck-revenue.csv',rows),('february-admin-audit.csv',admin),('excluded-historical-rows.csv',historical)]:
 with (out/filename).open('w') as f:
  writer=csv.DictWriter(f,fieldnames=list(values[0]));writer.writeheader();writer.writerows(values)
print(json.dumps({k:v for k,v in result.items() if k in ['totals','truck_count','nonzero_cent_differences','full_original_import_error','synthetic_calculation']},indent=2,default=str))
