"""Read-only monthly expense preview. No database, API or legacy payslip imports.

Persist HeaderMapping records at confirmation in a future phase; pass those records
back as approved_mappings on later imports. Unknown columns are never expenses
until explicitly mapped. None in approved_mappings explicitly ignores a column.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import io
import re
import unicodedata
from typing import Mapping
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

CENT = Decimal('0.01')
OPTIONAL_CATEGORY = 'kasarani_ksm_fuel'


def normalize_header(value: str) -> str:
    value = unicodedata.normalize('NFKC', value).casefold()
    return ' '.join(re.sub(r'[^\w]+', ' ', value).split())


KNOWN = {
    'fuel': ('Fuel', ('fuel', 'fuel cost')),
    'kasarani_ksm_fuel': ('Kasarani & KSM Fuel', ('kasarani & ksm fuel', 'kasarani and ksm fuel', 'kasarani ksm fuel')),
    'breakages': ('Breakages', ('breakages', 'breakage')),
    'telematics_vss': ('Telematics & VSS', ('telematics & vss', 'telematics and vss', 'telematics vss')),
    'gprs': ('GPRS', ('gprs',)),
    'shunting': ('Shunting', ('shunting',)),
}
ALIASES = {normalize_header(alias): key for key, (_, aliases) in KNOWN.items() for alias in aliases}
TRUCK_HEADERS = {normalize_header(s) for s in ('TRUCK REG/NO', 'TRUCK REG.NO', 'TRUCK REGISTRATION', 'TRUCK REG NO')}
IGNORED = {normalize_header(s) for s in ('OPERATING EXPENSES', 'TOTAL OPERATING EXPENSES', 'TOTAL EXPENSES', 'GROSS AMOUNT', 'NET PAID I', 'NET PAID 1', 'NET PAID 2', 'NET PAID', 'NET PAY', '7% COMMISSION', '7 COMMISSION')}


@dataclass(frozen=True)
class Issue:
    code: str
    message: str
    cell: str | None = None
    blocking: bool = True


@dataclass(frozen=True)
class HeaderMapping:
    column: int
    original_header: str
    normalized_header: str
    classification: str  # truck / expense / ignored / unknown
    category: str | None = None
    label: str | None = None
    approval_source: str = 'built_in'


@dataclass
class ExpenseRow:
    row_number: int
    original_plate: str
    plate: str
    expenses: dict[str, Decimal | None] = field(default_factory=dict)
    source_values: dict[str, dict] = field(default_factory=dict)
    issues: list[Issue] = field(default_factory=list)
    classifications_complete: bool = True

    @property
    def ready(self) -> bool:
        return self.classifications_complete and not any(i.blocking for i in self.issues)


@dataclass
class ExpensePreview:
    month: date
    sheet: str
    header_row: int
    headers: list[HeaderMapping]
    rows: list[ExpenseRow]
    issues: list[Issue]

    @property
    def ready(self) -> bool:
        return bool(self.rows) and not any(i.blocking for i in self.issues) and all(r.ready for r in self.rows)


def money(value: object) -> Decimal:
    """Validate and round an Excel amount; zero is distinct from missing."""
    if value is None or isinstance(value, bool) or isinstance(value, (date, datetime)):
        raise ValueError('Expected a numeric amount; blank is not zero')
    text = str(value).strip()
    if not isinstance(value, (int, float, Decimal)) and not re.fullmatch(r'[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?', text):
        raise ValueError('Invalid monetary amount')
    try:
        amount = Decimal(text.replace(',', ''))
        if not amount.is_finite() or abs(amount) >= Decimal('10000000000000000'):
            raise ValueError('Amount outside supported range')
        return amount.quantize(CENT, rounding=ROUND_HALF_UP)
    except InvalidOperation as exc:
        raise ValueError('Invalid monetary amount') from exc


def category_key(label: str) -> str:
    normalized = normalize_header(label)
    if normalized in IGNORED or normalized in TRUCK_HEADERS:
        raise ValueError('Financial summary and truck headers cannot be expense categories')
    key = ALIASES.get(normalized) or normalized.replace(' ', '_')
    if not key:
        raise ValueError('Expense category must have a name')
    return key


def validate_xlsx_envelope(content: bytes) -> None:
    """Shared read-only size/container guard for expense and raw-review uploads."""
    if len(content) > 20_000_000:
        raise ValueError('Workbook exceeds 20 MB')
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100_000_000:
                raise ValueError('Expanded workbook exceeds 100 MB')
            if not {'[Content_Types].xml', 'xl/workbook.xml', '_rels/.rels'}.issubset(archive.namelist()):
                raise ValueError('Expected an XLSX workbook structure')
            if any(name.lower().endswith('vbaproject.bin') for name in archive.namelist()):
                raise ValueError('Macro-enabled workbooks are not supported')
    except BadZipFile as exc:
        raise ValueError('Expected an XLSX workbook') from exc


def parse_expenses(
    content: bytes, month: date, *, sheet_name: str | None = None,
    approved_mappings: Mapping[str, str | None] | None = None,
    corrections: Mapping[str, object] | None = None,
) -> ExpensePreview:
    """Preview XLSX; mappings/corrections represent explicit administrator decisions."""
    if month.day != 1:
        raise ValueError('Reporting month must be its first day')
    validate_xlsx_envelope(content)
    wb = load_workbook(io.BytesIO(content), data_only=False)
    cached = load_workbook(io.BytesIO(content), data_only=True)
    try:
        if sheet_name is None and len(wb.sheetnames) != 1:
            raise ValueError('Select a worksheet explicitly for a multi-sheet workbook')
        sheet_name = sheet_name or wb.sheetnames[0]
        if sheet_name not in wb.sheetnames:
            raise ValueError('Worksheet not found')
        ws, values = wb[sheet_name], cached[sheet_name]
        if ws.max_row > 10000 or ws.max_column > 200:
            raise ValueError('Worksheet exceeds preview limits')
        header_rows = [r for r in range(1, min(ws.max_row, 30) + 1)
                       if any(normalize_header(str(c.value)) in TRUCK_HEADERS for c in ws[r] if c.value is not None)]
        if len(header_rows) != 1:
            raise ValueError('Expected one unambiguous truck header within first 30 rows')
        header_row = header_rows[0]
        issues: list[Issue] = []
        supplied = approved_mappings or {}
        mappings: dict[str, str | None] = {}
        for original, label in supplied.items():
            norm = normalize_header(original)
            if norm in mappings:
                raise ValueError('Duplicate normalized mapping header')
            if norm in IGNORED or norm in TRUCK_HEADERS or norm in ALIASES:
                raise ValueError('Built-in header classifications cannot be overridden')
            mappings[norm] = label
        headers: list[HeaderMapping] = []
        used_headers: set[str] = set()
        used_categories: set[str] = set()
        for cell in ws[header_row]:
            if cell.value is None:
                if any(ws.cell(r, cell.column).value is not None for r in range(header_row + 1, ws.max_row + 1)):
                    issues.append(Issue('missing_header', 'Populated column has no header', cell.coordinate))
                continue
            original = str(cell.value)
            norm = normalize_header(original)
            if norm in used_headers:
                issues.append(Issue('duplicate_header', 'Duplicate normalized header', cell.coordinate))
            used_headers.add(norm)
            category = label = None
            source = 'built_in'
            if norm in TRUCK_HEADERS:
                classification = 'truck'
            elif norm in IGNORED:
                classification = 'ignored'
            elif norm in ALIASES:
                classification = 'expense'
                category = ALIASES[norm]
                label = KNOWN[category][0]
            elif norm in mappings:
                source = 'administrator'
                label = mappings[norm]
                classification = 'ignored' if label is None else 'expense'
                category = None if label is None else category_key(label)
            else:
                classification = 'unknown'
                issues.append(Issue('unknown_column', 'Classify this column as expense or ignored', cell.coordinate))
            if category is not None:
                if category in used_categories:
                    issues.append(Issue('duplicate_category', f'Multiple columns map to {category}', cell.coordinate))
                used_categories.add(category)
            headers.append(HeaderMapping(cell.column, original, norm, classification, category, label, source))
        trucks = [h for h in headers if h.classification == 'truck']
        if len(trucks) != 1:
            raise ValueError('Expected exactly one truck registration column')
        if not used_categories:
            issues.append(Issue('no_expense_categories', 'At least one approved expense category is required'))
        # The optional absent category is the only automatically assumed zero.
        optional_absent = OPTIONAL_CATEGORY not in used_categories
        if optional_absent:
            issues.append(Issue('optional_category_absent', 'Kasarani & KSM Fuel column absent; assuming 0.00', blocking=False))
        # Explicitly compare recognizable workbook month evidence with selection.
        months = {name.casefold(): i for i, name in enumerate(
            ('January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December'), 1)}
        evidence = []
        for row in ws.iter_rows(min_row=1, max_row=max(1, header_row - 1)):
            for cell in row:
                for name, year in re.findall(r'\b(' + '|'.join(months) + r')\s+(20\d{2})\b', str(cell.value), re.I):
                    evidence.append(date(int(year), months[name.casefold()], 1))
        if any(value != month for value in evidence):
            issues.append(Issue('month_mismatch', 'Workbook title disagrees with selected month'))
        if not evidence:
            issues.append(Issue('month_unverified', 'No month evidence; explicit administrator month confirmation required'))
        for merged in ws.merged_cells.ranges:
            if merged.max_row >= header_row:
                issues.append(Issue('merged_data_cells', 'Merged cells intersect headers/data', str(merged)))
        rows: list[ExpenseRow] = []
        seen: set[str] = set()
        corrections = corrections or {}
        consumed: set[str] = set()
        for number in range(header_row + 1, ws.max_row + 1):
            cells = list(ws[number])
            if all(c.value is None or (isinstance(c.value, str) and not c.value.strip()) for c in cells):
                continue
            truck_cell = ws.cell(number, trucks[0].column)
            original_plate = '' if truck_cell.value is None else str(truck_cell.value)
            if normalize_header(original_plate) in {'total', 'grand total', 'totals'}:
                continue
            plate = ''.join(unicodedata.normalize('NFKC', original_plate).split()).upper()
            row = ExpenseRow(number, original_plate, plate, classifications_complete=not any(i.blocking for i in issues))
            if not re.fullmatch(r'[A-Z0-9]+', plate) or truck_cell.data_type == 'f':
                row.issues.append(Issue('invalid_truck', 'Missing or invalid truck registration', truck_cell.coordinate))
            if plate in seen:
                row.issues.append(Issue('duplicate_truck', 'Duplicate normalized registration', truck_cell.coordinate))
            seen.add(plate)
            for header in headers:
                cell = ws.cell(number, header.column)
                if header.classification in {'truck', 'ignored'}:
                    continue
                value = values.cell(number, header.column).value if cell.data_type == 'f' else cell.value
                source = {'cell': cell.coordinate, 'header': header.original_header,
                          'original_value': str(cell.value) if cell.value is not None else None,
                          'cached_value': str(value) if cell.data_type == 'f' and value is not None else None}
                row.source_values[cell.coordinate] = source
                if header.classification != 'expense':
                    continue
                if cell.coordinate in corrections:
                    value = corrections[cell.coordinate]
                    consumed.add(cell.coordinate)
                    source['approved_correction'] = str(value)
                elif cell.data_type == 'f':
                    row.issues.append(Issue('formula_cache_missing' if value is None else 'formula_cached',
                                            'Formula requires cached value; cached values may be stale', cell.coordinate, value is None))
                try:
                    amount = money(value)
                except ValueError:
                    amount = None
                    row.issues.append(Issue('blank_expense' if value is None or str(value).strip() == '' else 'invalid_expense',
                                            'Correct or explicitly confirm this cell as a numeric amount', cell.coordinate))
                if amount is not None and amount < 0:
                    row.issues.append(Issue('negative_expense', 'Negative deduction retained for review', cell.coordinate, False))
                # Do not overwrite a prior value when mappings conflict.
                if header.category not in row.expenses:
                    row.expenses[header.category] = amount
            if optional_absent:
                row.expenses[OPTIONAL_CATEGORY] = Decimal('0.00')
            rows.append(row)
        if set(corrections) != consumed:
            raise ValueError('Corrections must refer to included expense cells on truck rows')
        if not rows:
            issues.append(Issue('no_trucks', 'No truck expense rows found'))
        return ExpensePreview(month, sheet_name, header_row, headers, rows, issues)
    finally:
        wb.close()
        cached.close()
