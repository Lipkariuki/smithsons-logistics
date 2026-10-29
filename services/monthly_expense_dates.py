"""Read-only raw-source date/provenance review. Never imports or alters DHL rows."""
from __future__ import annotations
import base64
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
import hashlib
import io
import re
import unicodedata
from openpyxl import load_workbook
from services.monthly_expense_parser import validate_xlsx_envelope, normalize_header, money


def _plate(value):
    return ''.join(unicodedata.normalize('NFKC',str(value or '')).split()).upper()


def _ref(value):
    text = '' if value is None else str(value).strip()
    return text[:-2] if text.endswith('.0') else text


def inspect_source(content: bytes, month: date):
    validate_xlsx_envelope(content)
    wb = load_workbook(io.BytesIO(content), data_only=True)
    try:
        if len(wb.worksheets) != 1:
            raise ValueError('Raw review requires a single worksheet')
        ws = wb.active
        if ws.max_row > 10000 or ws.max_column > 200:
            raise ValueError('Raw review workbook exceeds preview limits')
        candidates = []
        for number in range(1, min(ws.max_row, 30) + 1):
            names = {normalize_header(str(c.value)): c.column for c in ws[number] if c.value is not None}
            if {'date', 'truck', 'ref no', 'cost', 'amount'}.issubset(names):
                candidates.append((number, names))
        if len(candidates) != 1:
            raise ValueError('Raw report needs one DATE/TRUCK/REF NO/COST/AMOUNT header')
        first, names = candidates[0]
        issues, groups = [], {}
        for number in range(first + 1, ws.max_row + 1):
            values = {key: ws.cell(number, col).value for key, col in names.items()}
            # A footer with no date, truck, reference or charge type is not a transaction.
            if all(values.get(k) is None for k in ('date', 'truck', 'ref no', 'cost')):
                continue
            value = values['date']
            day = value.date() if isinstance(value, datetime) else value if isinstance(value, date) else None
            if isinstance(value, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}', value.strip()):
                try: day = date.fromisoformat(value.strip())
                except ValueError: pass
            code = 'missing_source_date' if value is None else 'invalid_source_date' if day is None else 'out_of_period_date' if (day.year, day.month) != (month.year, month.month) else None
            if code:
                issues.append({'source_row': number, 'date': value.isoformat() if isinstance(value, (date, datetime)) else str(value) if value is not None else None,
                               'truck': str(values.get('truck') or ''), 'reference': _ref(values.get('ref no')),
                               'amount': str(values['amount']) if values.get('amount') is not None else None,
                               'code': code, 'reason': f'Expected {month:%Y-%m}; original date requires review'})
                continue
            plate, ref = _plate(values.get('truck')), _ref(values.get('ref no'))
            if not plate or not ref:
                issues.append({'source_row': number, 'date': day.isoformat(), 'truck': plate, 'reference': ref,
                               'amount': str(values.get('amount')), 'code': 'invalid_source_identity',
                               'reason': 'Dated source row needs a truck and reference'})
                continue
            key = (day.isoformat(), plate, ref)
            item = groups.setdefault(key, [Decimal('0.00'), Decimal('0.00')])
            label = str(values.get('cost') or '').strip().lower()
            # Same verified gross rule. Admin contributes zero; never an expense.
            if 'dist' in label or 'off' in label:
                try: amount = money(values.get('amount'))
                except ValueError:
                    issues.append({'source_row': number, 'date': day.isoformat(), 'truck': plate, 'reference': ref,
                                   'amount': str(values.get('amount')), 'code': 'invalid_source_amount',
                                   'reason': 'Transport amount is invalid; correct the source through an approved process'})
                    continue
                item[0 if 'dist' in label else 1] += amount
        return {'sha256': hashlib.sha256(content).hexdigest(), 'sheet': ws.title,
                'reporting_month': month.isoformat(), 'date_issues': issues,
                'out_of_month_records': [r for r in issues if r['code'] == 'out_of_period_date'],
                'unparsed_date_rows': [r['source_row'] for r in issues if r['code'] == 'invalid_source_date'],
                'action': 'review_only_no_source_changes'}, groups
    finally:
        wb.close()


def retain_source(content, month):
    review, _ = inspect_source(content, month)
    review['source_content_base64'] = base64.b64encode(content).decode('ascii')
    return review


def charge_issues(content, month, target_plate=None):
    """Flag ambiguous/repeated transport charges; never silently deduplicate.

    Keep these diagnostics separate from the immutable v1 review metadata so
    previously retained sources remain readable without rewriting their audit.
    """
    wb = load_workbook(io.BytesIO(content), data_only=True)
    try:
        ws = wb.active
        for number in range(1, min(ws.max_row, 30) + 1):
            names = {normalize_header(str(c.value)): c.column for c in ws[number] if c.value is not None}
            if {'date', 'truck', 'ref no', 'cost', 'amount'}.issubset(names):
                break
        else:
            return []  # inspect_source reports malformed headers first.
        found, result = {}, []
        for row in range(number + 1, ws.max_row + 1):
            value = ws.cell(row, names['date']).value
            day = value.date() if isinstance(value, datetime) else value if isinstance(value, date) else None
            if isinstance(value, str):
                try: day = date.fromisoformat(value.strip())
                except ValueError: pass
            if day is None or (day.year, day.month) != (month.year, month.month):
                continue  # Strict date validation handles these separately.
            plate = _plate(ws.cell(row, names['truck']).value)
            if target_plate and plate and plate != target_plate:
                continue
            label = str(ws.cell(row, names['cost']).value or '').strip().casefold()
            if label == 'admin':
                continue
            if label not in {'dist', 'offloading'}:
                result.append({'code': 'unsupported_source_charge', 'cell': None, 'blocking': True,
                    'message': f'Excel row {row}: COST must be DIST, Offloading or Admin. Review the source charge type.'})
                continue
            try: amount = money(ws.cell(row, names['amount']).value)
            except ValueError: continue  # inspect_source already reports it.
            key = (day, plate, _ref(ws.cell(row, names['ref no']).value), label, amount)
            if key in found:
                result.append({'code': 'duplicate_source_row', 'cell': None, 'blocking': True,
                    'message': f'Excel rows {found[key]} and {row}: repeated {label} charge for {plate}, ref {key[2]}. Check the original; no duplicate has been removed automatically.'})
            else:
                found[key] = row
        return result
    finally:
        wb.close()


def validate_source(review, dispositions, source, month, *, target_plate=None, _inspected=None):
    """Reparse retained bytes every build/confirmation; flags cannot supply results."""
    def issue(code, message): return {'code': code, 'message': message, 'cell': None, 'blocking': True}
    if not review or not review.get('source_content_base64'):
        return [issue('source_dates_unreviewed', 'Attach the original raw trips workbook and review its dates before confirming.')]
    try:
        content = base64.b64decode(review['source_content_base64'], validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError('Retained raw workbook is invalid') from exc
    actual, expected = _inspected if _inspected is not None else inspect_source(content, month)
    if {k: v for k, v in review.items() if k != 'source_content_base64'} != actual:
        raise ValueError('Raw source review checksum or reporting month changed')
    known = {str(r['source_row']): r for r in actual['date_issues']}
    if set(dispositions) - set(known):
        raise ValueError('Date decisions must identify affected rows in this exact raw workbook')
    result = charge_issues(content, month, target_plate)
    for key, row in known.items():
        decision = dispositions.get(key)
        # Only out-of-period dates can be excluded through this narrow workflow.
        # Missing/invalid dates or amounts require source correction, never guessing.
        if decision and (row['code'] != 'out_of_period_date' or decision.get('action') != 'exclude_from_period' or not decision.get('reason', '').strip()):
            raise ValueError('Only historical dates can be explicitly excluded, with a review reason')
        related = not target_plate or not re.fullmatch(r'[A-Z0-9]+', _plate(row['truck'])) or _plate(row['truck']) == target_plate
        if not decision and related:
            result.append(issue(row['code'], f"Excel row {key}: {row['reason']}"))
    observed = {}
    for row in source:
        if target_plate and _plate(row['truck_plate']) != target_plate:
            continue
        key = (row['date'], _plate(row['truck_plate']), _ref(row['ref_no']))
        item = observed.setdefault(key, [Decimal('0.00'), Decimal('0.00')])
        try:
            for index, field in enumerate(('distribution_cost', 'offloading_cost')):
                item[index] += money(row[field])
            if money(row['total_revenue']) != money(row['distribution_cost']) + money(row['offloading_cost']):
                result.append(issue('source_gross_inconsistent', 'Stored DHL gross differs from its transport charges; source review required.'))
        except ValueError:
            result.append(issue('invalid_source_amount','Stored transport amount is invalid; review its source.'))
    if target_plate:
        expected = {key: value for key, value in expected.items() if key[1] == target_plate}
    if expected != observed:
        result.append(issue('source_manifest_mismatch', 'The retained raw workbook does not match selected-month DHL records. Check source completeness; no revenue has been changed.'))
    return result
