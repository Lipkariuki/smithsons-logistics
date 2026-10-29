"""Snapshot-only API presentation, manual comparisons and branded exports."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
import io
import textwrap

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from services.monthly_expense_reconciliation import reference_analysis
from utils.payslip_branding import payslip_text

FINANCIAL = ('gross', 'operating_expenses', 'commission_base', 'commission', 'net_payable')
LABELS = ('Gross Transport Revenue', 'Operating Expenses', 'Commission Base', '7% Commission', 'Net Payable')
REFERENCE_LABELS = {'gross': 'Gross Revenue', 'commission': 'Commission', 'net_payable': 'Net Payable', 'operating_expenses': 'Operating Expenses'}
WARNING_LABELS = {
    'optional_category_absent': 'Optional fuel category absent',
    'no_matching_transport_revenue_verified_zero': 'Verified no business',
    'zero_revenue': 'Zero revenue',
    'negative_net_balance_review': 'Negative balance: review required',
    'raw_adjustments_not_reviewed': 'Historical adjustments not reviewed',
}
WARNING_MESSAGES = {
    'no_matching_transport_revenue_verified_zero': 'No monthly transport business was explicitly verified; expenses are retained.',
    'zero_revenue': 'Gross revenue is KES 0.00; all approved expenses remain included.',
    'negative_net_balance_review': 'Expenses exceed gross revenue. Review the negative balance; it has not been written off.',
}


def totals(rows):
    """Unknown values stay unknown; disclose complete-row subtotals separately."""
    output = {}
    for key in FINANCIAL:
        values = [r[key] for r in rows]
        complete = all(v is not None for v in values)
        output[key] = {'amount': sum(values, Decimal('0.00')) if complete else None,
                       'known_subtotal': sum((v for v in values if v is not None), Decimal('0.00')),
                       'missing_rows': sum(v is None for v in values)}
    return output


def report_view(batch, versions=(), *, owner_id=None, vehicle_id=None):
    references = reference_analysis(batch)
    detail = {k: v for k, v in batch.items() if k not in {'source_content', 'rows'}}
    selected = [dict(r) for r in batch['rows'] if
                (owner_id is None or r['vehicle_owner_snapshot'].get('owner_id') == owner_id)
                and (vehicle_id is None or r['vehicle_id'] == vehicle_id)]
    for row in selected:
        row['comparisons'] = references[row['source_row']]
    detail.update(rows=selected, versions=list(versions), totals=totals(selected),
                  full_report_totals=totals(batch['rows']),
                  filter_scope={'owner_id': owner_id, 'vehicle_id': vehicle_id},
                  ready=bool(batch['rows']) and not any(i['blocking'] for i in batch['issues'])
                  and not any(i['blocking'] for r in batch['rows'] for i in r['issues']),
                  comparison_policy={'difference': 'calculated_minus_manual', 'rounding_sized_tolerance': '0.01',
                                     'version': 'manual-reference-v1'})
    if batch['status'] == 'draft' and not any((batch['decisions'].get('historical_review') or {}).get(k) for k in ('source_content_base64', 'archive_id')):
        detail['ready'] = False
        detail['issues'] = list(detail['issues'])
        if not any(i['code'] == 'source_dates_unreviewed' for i in detail['issues']):
            detail['issues'].append({'code':'source_dates_unreviewed','message':'Attach and review the retained raw source before confirming.','blocking':True,'cell':None})
    detail['can_confirm'] = batch['status'] == 'draft' and detail['ready']
    # Original raw workbook is retained server-side only, not returned in JSON.
    detail['decisions'] = dict(detail['decisions'])
    if detail['decisions'].get('historical_review'):
        detail['decisions']['historical_review'] = {k: v for k, v in detail['decisions']['historical_review'].items() if k != 'source_content_base64'}
    return decimal_json(detail)


def decimal_json(value):
    if isinstance(value, Decimal):
        return format(value, '.2f')
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return {k: decimal_json(v) for k, v in value.items()}
    if isinstance(value, (tuple, list)):
        return [decimal_json(v) for v in value]
    return value


def warnings(view):
    output = [('Report', issue['code'], issue['message']) for issue in view['issues']]
    for row in view['rows']:
        output += [(row['normalized_plate'], i['code'], WARNING_MESSAGES.get(i['code'], i['message'])) for i in row['issues']]
        for key, comparison in row['comparisons'].items():
            if comparison.get('cached_formula'):
                output.append((row['normalized_plate'], 'reference_formula_cached',
                               f'{REFERENCE_LABELS[key]}: manual reference uses a cached formula value; it may be stale'))
            if comparison['status'] != 'exact_match':
                output.append((row['normalized_plate'], comparison['status'],
                               f"{REFERENCE_LABELS[key]}: calculated minus manual {comparison.get('difference') or 'unavailable'} KES"))
    review = view['decisions'].get('historical_review')
    if review:
        for record in review.get('date_issues', review.get('out_of_month_records', [])):
            code = 'historical_adjustment_review' if record.get('code', 'out_of_period_date') == 'out_of_period_date' else record['code']
            decision = view['decisions'].get('date_dispositions', {}).get(str(record['source_row']), {})
            output.append((record['truck'], code,
                           f"Source row {record['source_row']}, date {record['date']}, truck {record['truck']}, reference {record['reference']}, amount {record['amount']}; review decision: {decision.get('reason', 'PENDING')}"))
    else:
        output.append(('Report', 'raw_adjustments_not_reviewed',
                       'Historical adjustments have not been reviewed against the raw transport report.'))
    return output


def _text(value):
    return payslip_text(str(value)).replace('\x00', '')


def _excel_text(value):
    text = _text(value)
    # Excel must never execute workbook labels, owner names or administrator notes.
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) else text


def export_xlsx(view):
    wb = Workbook()
    wb.properties.creator = 'SMITHSONS LOGISTICS'
    wb.properties.title = 'Monthly Transport Statement'
    wb.properties.subject = 'Monthly Transport Statement'
    ws = wb.active
    ws.title = 'Monthly Summary'
    ws.append(['SMITHSONS LOGISTICS'])
    ws.append(['Monthly Transport Statement', view['reporting_month'], 'Version', view['version'], 'Status', view['status']])
    ws.append(['Currency', 'KES'])
    if view['status'] != 'confirmed':
        ws.append(['REVIEW COPY — ' + view['status'].upper()])
    categories = {}
    for row in view['rows']:
        for key, expense in row['deductions'].items():
            categories.setdefault(key, expense['label'])
    headers = ['Truck', 'Owner', LABELS[0], *categories.values(), *LABELS[1:]]
    difference_columns = set()
    for key in tuple(REFERENCE_LABELS):
        label = REFERENCE_LABELS[key]
        difference_columns.add(len(headers) + 2)
        headers += [f'Manual {label}', f'Difference {label} (calculated minus manual)', f'Comparison {label}']
    ws.append([_excel_text(h) for h in headers])
    header_row = ws.max_row
    for row in view['rows']:
        data = [_excel_text(row['normalized_plate']), _excel_text(row['vehicle_owner_snapshot'].get('owner_name') or 'Unassigned'),
                _number(row['gross'])]
        data += [_number(row['deductions'].get(key, {}).get('amount')) for key in categories]
        data += [_number(row[key]) for key in FINANCIAL[1:]]
        for key in tuple(REFERENCE_LABELS):
            comparison = row['comparisons'][key]
            data += [_number(comparison.get('manual')), _number(comparison.get('difference')),
                     _excel_text(comparison['status'].replace('_', ' ').capitalize())]
        ws.append(data)
    # Do not silently treat missing row/category amounts as zero.
    category_totals = []
    for key in categories:
        values = [r['deductions'].get(key, {}).get('amount') for r in view['rows']]
        category_totals.append(None if any(v is None for v in values) else sum((_number(v) for v in values), Decimal('0.00')))
    ws.append(['MONTHLY TOTAL', None, _number(view['totals']['gross']['amount']), *category_totals,
               *[_number(view['totals'][key]['amount']) for key in FINANCIAL[1:]]])
    total_row = ws.max_row
    ws.append(['Commission is calculated separately for each truck, then totalled.'])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=len(headers))
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[ws.max_row].height = 30
    ws.freeze_panes = f'C{header_row + 1}'
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{total_row - 1}'
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='352166')
        cell.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[header_row].height = 60
    for cell in ws[total_row]:
        cell.font = Font(bold=True, color='352166')
        cell.fill = PatternFill('solid', fgColor='EEE9F7')
    ws['A1'].font = Font(bold=True, size=18, color='352166')
    ws.merge_cells('A1:D1')
    ws['A2'].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 34
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 24 if i > 2 else 22
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            if isinstance(cell.value, Decimal):
                cell.number_format = ('+#,##0.00;[Red]-#,##0.00;+0.00' if cell.column in difference_columns
                                      else '#,##0.00;[Red](#,##0.00)')
    notes = wb.create_sheet('Review Warnings')
    notes.append(['SMITHSONS LOGISTICS'])
    notes.append(['Truck / Scope', 'Issue', 'Details'])
    for warning in warnings(view):
        scope, code, message = warning
        notes.append([_excel_text(scope), _excel_text(WARNING_LABELS.get(code, code.replace('_', ' ').capitalize())), _excel_text(message)])
    notes.column_dimensions['A'].width = 24
    notes.column_dimensions['B'].width = 38
    notes.column_dimensions['C'].width = 110
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        lines = max(len(textwrap.wrap(str(cell.value or ''), width=width)) or 1
                    for cell, width in zip(row, (24, 38, 110)))
        notes.row_dimensions[row[0].row].height = 15 * lines + 8
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def _number(value):
    return None if value is None else Decimal(value)


def export_pdf(view):
    class StatementPDF(FPDF):
        def header(self):
            self.set_font('Helvetica', 'B', 18)
            self.cell(0, 10, 'SMITHSONS LOGISTICS', new_x='LMARGIN', new_y='NEXT', align='C')
            self.set_font('Helvetica', '', 11)
            self.cell(0, 7, 'Monthly Transport Statement', new_x='LMARGIN', new_y='NEXT', align='C')
            self.ln(4)

        def footer(self):
            self.set_y(-15)
            self.set_font('Helvetica', '', 9)
            self.cell(0, 8, f'SMITHSONS LOGISTICS | Page {self.page_no()}', align='C')

    pdf = StatementPDF()
    pdf.set_title('Monthly Transport Statement')
    pdf.set_author('SMITHSONS LOGISTICS')
    pdf.set_creator('SMITHSONS LOGISTICS')
    pdf.set_auto_page_break(True, margin=20)

    def line(value, bold=False):
        pdf.set_font('Helvetica', 'B' if bold else '', 10)
        pdf.multi_cell(0, 6, _text(value).encode('latin1', 'replace').decode('latin1'), new_x='LMARGIN', new_y='NEXT')

    def amount(label, value):
        line(f'{label}: ' + ('Unavailable' if value is None else f'KES {Decimal(value):,.2f}'))

    pdf.add_page()
    line(f"Reporting month: {view['reporting_month'][:7]}", True)
    line(f"Version: {view['version']} | Status: {view['status'].upper()}", True)
    if view['status'] != 'confirmed':
        line('REVIEW COPY - this version is not the current confirmed report', True)
    line('Monthly totals', True)
    for key, label in zip(FINANCIAL, LABELS):
        amount(label, view['totals'][key]['amount'])
        if view['totals'][key]['missing_rows']:
            line(f"{view['totals'][key]['missing_rows']} row(s) unresolved; no complete monthly total.")
    line('Review warnings', True)
    line('Commission is calculated separately for each truck, then totalled.')
    for scope, code, message in warnings(view):
        line(f'{scope}: {WARNING_LABELS.get(code, code.replace("_", " "))} - {message}')
    for row in view['rows']:
        pdf.add_page()
        line(f"Truck: {row['normalized_plate']}", True)
        line(f"Owner: {row['vehicle_owner_snapshot'].get('owner_name') or 'Unassigned'}")
        line(f"Month: {view['reporting_month'][:7]} | Version: {view['version']} | Status: {view['status']}")
        amount(LABELS[0], row['gross'])
        line('Itemized Operating Expenses', True)
        for value in row['deductions'].values():
            amount(value['label'], value['amount'])
        for key, label in zip(FINANCIAL[1:], LABELS[1:]):
            amount(label, row[key])
        line('Manual Reference Comparison', True)
        for key, comparison in row['comparisons'].items():
            line(f"{REFERENCE_LABELS[key]}: {comparison['status'].replace('_', ' ')}; calculated minus manual "
                 f"{comparison.get('difference') or 'unavailable'} KES")
    return bytes(pdf.output())
